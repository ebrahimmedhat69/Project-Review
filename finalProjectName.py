import pandas as pd
import re
from sklearn.model_selection import train_test_split
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import classification_report

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ==========================
# 1- Read training sheet
# ==========================
file_path = "C:/Users/dell/OneDrive/Desktop/ML/31-12-2025.xlsx"

df = pd.read_excel(file_path)

# ==========================
# 2- Data cleaning
# ==========================

df['Description'] = df['Description'].fillna('')
df['ProductDescription'] = df['ProductDescription'].fillna('')

df['Maintenance job type'] = df['Maintenance job type'].fillna('')
df['Category'] = df['Category'].fillna('')
df['Project Name'] = df['Project Name'].fillna('unknown')

# merge the cols 
df['text'] = (
    df['Description'].astype(str) + " " +
    df['ProductDescription'].astype(str) + " " +
    df['Maintenance job type'].astype(str) + " " +
    df['Category'].astype(str)
)

X = df['text']
y = df['Project Name']

# ==========================
# 3-Train/Test split
# ==========================
X_train, X_test, y_train, y_test = train_test_split(
    X, y, test_size=0.2, random_state=42
)

# ==========================
# 4- TF-IDF vectorizer
# ==========================
vectorizer = TfidfVectorizer(
    ngram_range=(1, 2),
    min_df=2,
    max_df=0.9
)
X_train_vec = vectorizer.fit_transform(X_train)
X_test_vec = vectorizer.transform(X_test)

# ==========================
# 5- Logistic Regression model
# ==========================
model = LogisticRegression(
    max_iter=1000,
    class_weight='balanced'
)
model.fit(X_train_vec, y_train)

# ==========================
# 6- ML evaluation
# ==========================
y_pred = model.predict(X_test_vec)
print("=== ML REPORT ===")
print(classification_report(y_test, y_pred))

# ==========================
# 7- Fixed keywords rules
# ==========================

fixed_keywords_projects = {
    
    'Camp':['Maintancei work in camp'],

    'laboratory': [
        'xray lab',
        'laboratory requirements',
        'chemical requirements',
        'maintenance lab',
        'lab'
    ],
    
    'Electrical Regular Maintenance': [
        'Lamp Led',
        'Contactor',
        'Battery',
        'Capacitor',
        'install low voltage',
        'change cooling fan',
        'motion detector',
        'switch',
        'mcb',
        'socket',
        'change motor',
        'change auxiliary'
    ],

    
    

    'Mechanical Regular Maintenance': [
        'Anti rust spray',
        'Cleaning Cloth, كهن او اسطبة قماش',
        'replace drag chain',
        'change v belts',
        'v belt',
        'cutting disc',
        'washer',
        'nut',
        'bolt'
    ],


    'Environment': [
        'support cage for filter bags',
        'change bags',
        'wire cage',
        'bag filter',
        'filter bag',
        'pulse valve',
        'asco joucomatic',
        'repair kit',
        'impeller for fan',
        'Repair Kit, 2 Way Valve'
    ],

    'Tools': [
        'Mechanical Dep-Consum (Workshop Tools & Supplies Safety and'
    ],


    'Belt Conveyors': [
        'change return rollers',
        'rubber disc idler',
        'self adjusting idler',
        'impact roller',
        'side skirt',
        'belt',
        'Return Idler',
        'Idler,  Return',
        'Idler, Impact',
        'Idler, Self-aligning',
        'Idler, Common',
        'Idler, Steel',
        'Guide Idler',
        'Impact idler ',
        'Idler, Flat return',
        'Roller, Return',
        'Roller, Guide',
        'Roller, Steel',
        'Roller, Coated Rubber',
        'Impact Roller',
        'Belt Guiding Roller',
        'lagging'
    ],

    'Lubricants': [
        'cat hydo ced 30',
        'mobil shc',
        'oil rs ultra',
        'berulit ga',
        'berucoat af 438',
        'lubricant',
        'omala',
        'total rubia'
    ],

    'Compressors-Service kit (2000,4000,8000 hrs) & blowers': [
        'Oil Filling Plug',
        'Oil level Plug',
        'scheduled service for the compressor',
        'repair the blower',
        'overhaul for the blower',
        'sealing the blower',
        'service kit'
    ],

    'Overhaul': [
        'castable',
        'bricks'
    ],

    'Safety': [
        'fabricate safety ',
        'Fabricate Safe guard',
        'Fabricate safety ',
        'safety helmet',
        'clear goggles',
        'dust mask',
        'disposable cover',
        'porter gloves',
        'fabricate hand rail',
        'safety supplies'
    ]
    
}

for proj in fixed_keywords_projects:
    fixed_keywords_projects[proj] = sorted(
        fixed_keywords_projects[proj],
        key=len,
        reverse=True
    )

# ==========================
# 9- Text cleaning function
# ==========================
def clean_text(text):
    text = str(text).lower()
    text = re.sub(r'[^a-z0-9\s]', ' ', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

# ==========================
# 10- Rule-based prediction
# ==========================
def check_fixed_keywords(text):
    text_clean = clean_text(text)
    for proj, keywords in fixed_keywords_projects.items():
        for kw in keywords:
            kw_clean = clean_text(kw)
            if kw_clean in text_clean:
                confidence = min(0.95, 0.3 + len(kw_clean) / 50)
                return proj, confidence, kw
    return None, None, None

# ==========================
# 11- Hybrid prediction function
# ==========================
def predict_project(description, item, maintenance='', category=''):
    full_text = (
        str(description) + " " +
        str(item) + " " +
        str(maintenance) + " " +
        str(category)
    )

    rule_proj, rule_conf, rule_kw = check_fixed_keywords(full_text)
    if rule_proj:
        return rule_proj, rule_conf, 'Rule', rule_kw

    text = clean_text(full_text)
    vec = vectorizer.transform([text])
    probs = model.predict_proba(vec)[0]
    idx = probs.argmax()

    return (
        model.classes_[idx],
        probs[idx],
        'ML',
        'model prediction'
    )


# ==========================
# 13- New sheet prediction
# ==========================
new_file_path = "C:/Users/dell/OneDrive/Desktop/ML/new_data.xlsx"
new_df = pd.read_excel(new_file_path)

new_df['Description'] = new_df['Description'].fillna('')
new_df['ProductDescription'] = new_df['ProductDescription'].fillna('')
new_df['Maintenance job type'] = new_df['Maintenance job type'].fillna('')
new_df['Category'] = new_df['Category'].fillna('')

projects = []
confidences = []
methods = []
reasons = []

for idx, row in new_df.iterrows():
    proj, conf, method, reason = predict_project(
        row['Description'],
        row['ProductDescription'],
        row['Maintenance job type'],
        row['Category']
    )
    projects.append(proj)
    confidences.append(round(conf, 2))
    methods.append(method)
    reasons.append(reason)

new_df['Predicted Project Name'] = projects
new_df['Confidence'] = confidences
new_df['Method'] = methods
new_df['Reason'] = reasons

new_df['Status'] = new_df.apply(
    lambda x: 'MATCH' if x['Project Name'] == x['Predicted Project Name'] else 'MISMATCH',
    axis=1
)

output_path = "C:/Users/dell/OneDrive/Desktop/ML/new_predictions.xlsx"
new_df.to_excel(output_path, index=False)

# ==========================
# Excel formatting
# ==========================
wb = load_workbook(output_path)
ws = wb.active

end_row = ws.max_row
end_col = ws.max_column
table_ref = f"A1:{get_column_letter(end_col)}{end_row}"

table = Table(displayName="PredictionsTable", ref=table_ref)
table.tableStyleInfo = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=False
)
ws.add_table(table)

headers = [c.value for c in ws[1]]

proj_col = headers.index("Project Name")
pred_col = headers.index("Predicted Project Name")
status_col = headers.index("Status")
method_col = headers.index("Method")

for row in ws.iter_rows(min_row=2, max_row=end_row):
    if row[proj_col].value == row[pred_col].value:
        fill = PatternFill("solid", fgColor="C6EFCE")
        row[status_col].value = "MATCH"
    else:
        fill = PatternFill("solid", fgColor="FFC7CE")
        row[status_col].value = "MISMATCH"

    for idx in [proj_col, pred_col, status_col]:
        row[idx].fill = fill
        row[idx].font = Font(bold=True)

    if row[method_col].value == "Rule":
        row[method_col].fill = PatternFill("solid", fgColor="C6EFCE")
    else:
        row[method_col].fill = PatternFill("solid", fgColor="FFF2CC")

# ==========================
# Auto column width (larger)
# ==========================
for column in ws.columns:
    max_length = 0
    col_letter = column[0].column_letter

    for cell in column:
        if cell.value:
            max_length = max(max_length, len(str(cell.value)))

    ws.column_dimensions[col_letter].width = min(max_length + 6, 50)

wb.save(output_path)
print("DONE  File formatted and readable")
